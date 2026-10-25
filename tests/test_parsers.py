import csv

from app.rag.parsers import get_file_type, is_supported_file, iter_records, normalize_text, parse_file


def test_supported_types():
    assert is_supported_file("x.PDF") and is_supported_file("a/b/log.csv") and is_supported_file("bundle.tar.gz")
    assert not is_supported_file("virus.exe")
    assert get_file_type("w.xlsx") == "xlsx" and get_file_type("w.tsv") == "csv"


def test_normalize_text_superscripts():
    assert normalize_text("EQL40200D³ ½ in") == "EQL40200D3 1/2 in"


def test_iter_records_streams_csv_with_semicolons(tmp_path):
    p = tmp_path / "alarms.csv"
    with open(p, "w", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Code", "Text"])
        w.writerow(["SRVO-050", "Collision detect"])
        w.writerow(["", ""])
        w.writerow(["PNT1-128", "Bell speed deviation"])
    rows = list(iter_records(str(p)))
    assert [r[2][0] for r in rows] == ["SRVO-050", "PNT1-128"]
    assert rows[0][1] == ["Code", "Text"]


def test_iter_records_xlsx(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tags"
    ws.append(["Tag", "HiHi"])
    ws.append(["OVN2_Z3_TEMP_PV", 241.3])
    p = tmp_path / "tags.xlsx"
    wb.save(p)
    rows = list(iter_records(str(p)))
    assert rows == [("Tags", ["Tag", "HiHi"], ["OVN2_Z3_TEMP_PV", "241.3"], 2)]


def test_markdown_and_docx(tmp_path):
    md = tmp_path / "m.md"
    md.write_text("# Title\n\nSome text about SRVO-050.")
    r = parse_file(str(md))
    assert r.success and "SRVO-050" in r.pages[0].text

    import docx
    d = docx.Document()
    d.add_heading("Bell cup cleaning", level=2)
    d.add_paragraph("Remove the bell cup.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Part", "Torque"
    t.cell(1, 0).text, t.cell(1, 1).text = "Bell cup", "12 Nm"
    p = tmp_path / "p.docx"
    d.save(p)
    r = parse_file(str(p))
    assert r.success
    assert "### Bell cup cleaning" in r.pages[0].text
    table = r.pages[0].rich_tables[0] if r.pages[0].rich_tables else None
    assert table is not None and table.data_rows[0][1].text == "12 Nm"


def test_missing_file_is_reported_not_raised():
    r = parse_file("/nonexistent/file.pdf")
    assert not r.success and "not found" in r.error
